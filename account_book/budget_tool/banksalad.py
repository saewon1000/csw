"""뱅크샐러드 내보내기(xlsx) 파서.

각 파일에는 두 개의 시트가 있다.
- "가계부 내역": 거래 내역 (날짜, 시간, 타입, 대분류, 소분류, 내용, 금액, 화폐, 결제수단, 메모)
- "뱅샐현황": 재무현황(자산/부채), 투자현황 등 요약 정보
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd

TXN_SHEET = "가계부 내역"
STATE_SHEET = "뱅샐현황"

# 뱅샐현황 재무현황 자산 항목명 -> (그룹, 표시명)
ASSET_GROUPS = {
    "자유입출금 자산": ("저축", "자유입출금"),
    "신탁 자산": ("저축", "신탁"),
    "현금 자산": ("저축", "현금"),
    "저축성 자산": ("저축", "저축성"),
    "전자금융 자산": ("저축", "전자금융"),
    "기타 실물 자산": ("저축", "기타 실물"),
    "보험 자산": ("저축", "보험"),
    "연금 자산": ("저축", "연금"),
    "투자성 자산": ("투자", "주식"),
    "부동산": ("투자", "부동산"),
    "동산": ("투자", "동산"),
}
LIABILITY_LABELS = {"장기대출", "단기대출"}


@dataclass
class BankSaladData:
    user: str
    transactions: pd.DataFrame
    assets: pd.DataFrame          # columns: group, item, name, amount
    liabilities: pd.DataFrame     # columns: item, name, amount
    investments: pd.DataFrame     # columns: kind, broker, name, principal, value, ret
    source_file: str = ""
    warnings: list[str] = field(default_factory=list)


def _cell(ws, r, c):
    return ws.cell(row=r, column=c).value


def parse_transactions(ws, user: str) -> pd.DataFrame:
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if not row or row[0] is None:
            continue
        date, time, typ, cat, sub, content, amount, cur, pay, memo = (
            list(row) + [None] * 10
        )[:10]
        rows.append(
            {
                "user": user,
                "date": pd.to_datetime(date) if date is not None else pd.NaT,
                "time": str(time) if time is not None else "",
                "type": typ,
                "category": cat,
                "subcategory": sub,
                "content": content,
                "amount": float(amount) if amount is not None else 0.0,
                "pay": pay,
                "memo": memo,
            }
        )
    df = pd.DataFrame(rows)
    if not df.empty:
        df["year"] = df["date"].dt.year
        df["month"] = df["date"].dt.month
    return df


def parse_financial_state(ws):
    """뱅샐현황 재무현황 섹션에서 자산/부채/투자 항목을 추출."""
    max_row = ws.max_row
    start = None
    end = None
    inv_start = None
    for r in range(1, max_row + 1):
        v = _cell(ws, r, 2)  # col B
        if isinstance(v, str):
            if v.startswith("3.재무현황"):
                start = r
            elif (v.startswith("총자산") or v.startswith("순자산")) and start and not end:
                end = r
            elif v.startswith("5.투자현황"):
                inv_start = r
    assets, liabilities = [], []
    if start:
        stop = end or max_row
        current_group = None
        for r in range(start, stop + 1):
            b = _cell(ws, r, 2)   # 자산 항목
            c = _cell(ws, r, 3)   # 자산 상품명
            e = _cell(ws, r, 5)   # 자산 금액
            f = _cell(ws, r, 6)   # 부채 항목
            g = _cell(ws, r, 7)   # 부채 상품명
            i = _cell(ws, r, 9)   # 부채 금액
            if isinstance(b, str) and b in ASSET_GROUPS:
                current_group = b
            if current_group and c is not None:
                grp, disp = ASSET_GROUPS[current_group]
                assets.append(
                    {"group": grp, "item": disp, "name": c, "amount": float(e or 0)}
                )
            if isinstance(f, str) and f in LIABILITY_LABELS and g is not None:
                liabilities.append({"item": f, "name": g, "amount": float(i or 0)})
    # 투자현황 (section 5)
    investments = []
    if inv_start:
        for r in range(inv_start + 3, max_row + 1):
            b = _cell(ws, r, 2)
            if isinstance(b, str) and (b.startswith("총계") or b.startswith("6.")):
                break
            name = _cell(ws, r, 4)
            if name is None:
                continue
            investments.append(
                {
                    "kind": b,
                    "broker": _cell(ws, r, 3),
                    "name": name,
                    "principal": float(_cell(ws, r, 6) or 0),
                    "value": float(_cell(ws, r, 7) or 0),
                    "ret": float(_cell(ws, r, 8) or 0),
                }
            )
    return (
        pd.DataFrame(assets, columns=["group", "item", "name", "amount"]),
        pd.DataFrame(liabilities, columns=["item", "name", "amount"]),
        pd.DataFrame(
            investments,
            columns=["kind", "broker", "name", "principal", "value", "ret"],
        ),
    )


def load_file(path: str, user: str) -> BankSaladData:
    p = Path(path)
    warnings = []
    if not p.exists():
        raise FileNotFoundError(f"파일을 찾을 수 없습니다: {path}")
    wb = openpyxl.load_workbook(p, data_only=True)
    if TXN_SHEET in wb.sheetnames:
        txns = parse_transactions(wb[TXN_SHEET], user)
    else:
        txns = pd.DataFrame()
        warnings.append(f"'{TXN_SHEET}' 시트가 없습니다.")
    if STATE_SHEET in wb.sheetnames:
        assets, liabilities, investments = parse_financial_state(wb[STATE_SHEET])
    else:
        assets = pd.DataFrame(columns=["group", "item", "name", "amount"])
        liabilities = pd.DataFrame(columns=["item", "name", "amount"])
        investments = pd.DataFrame(
            columns=["kind", "broker", "name", "principal", "value", "ret"]
        )
        warnings.append(f"'{STATE_SHEET}' 시트가 없습니다.")
    wb.close()
    return BankSaladData(
        user=user,
        transactions=txns,
        assets=assets,
        liabilities=liabilities,
        investments=investments,
        source_file=str(p),
        warnings=warnings,
    )
